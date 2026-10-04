"""Actual isolated MySQL/SQLite, original Tool, HTTP and provider-wire P5 acceptance."""
import asyncio
import base64
import io
import json
import os
from pathlib import Path
import zipfile
import httpx
import pytest
from sqlalchemy import text
from test_integration import env,client,query
from test_p5_delivery import TREE

ADMIN='/api/v1/admin/'

def login(env,client):
    with env[3].begin() as c:
        c.execute(text("INSERT INTO admin_user(user_id,username,password,status) VALUES('p5-admin','p5-admin','isolated-password',1) ON DUPLICATE KEY UPDATE status=1"))
    result=client.post(ADMIN+'admin-user/login',json={'username':'p5-admin','password':'isolated-password'})
    assert result.json()['code']=='0000',result.text
    assert 'HttpOnly' in result.headers['set-cookie'] and 'SameSite=lax' in result.headers['set-cookie']
    return result.json()['data']

def test_p5_admin_configuration_full_crud_auth_and_hot_effect(env,client):
    login(env,client)
    defs={
     'ai-client-api':dict(apiId='p5-api',baseUrl=os.environ['P1_TEST_MODEL_URL'],apiKey='isolated-provider-key',completionsPath='/v1/chat/completions',embeddingsPath='/v1/embeddings',status=1),
     'ai-client-model':dict(modelId='p5-model',apiId='p5-api',modelName='p5-actual-binding',modelType='chat',modelUsage='primary',supportsThinking=1,contextWindow=65536,status=1),
     'ai-client-tool-mcp':dict(mcpId='p5-mcp',mcpName='P5_DISABLED',transportType='stdio',transportConfig=json.dumps({'stdio':{'P5_DISABLED':{'command':'false'}}}),requestTimeout=2,status=0),
     'admin-user':dict(userId='p5-secondary',username='secondary',password='isolated-other-password',status=1)}
    for kind,data in defs.items():
        value=client.post(ADMIN+kind+'/create',json=data).json();assert value['code']=='0000' and value['data'] is True,(kind,value)
        rows=client.get(ADMIN+kind+'/query-all').json()['data'];record=next(r for r in rows if r[{'admin-user':'userId','ai-client-api':'apiId','ai-client-model':'modelId','ai-client-tool-mcp':'mcpId'}[kind]]==data[{'admin-user':'userId','ai-client-api':'apiId','ai-client-model':'modelId','ai-client-tool-mcp':'mcpId'}[kind]])
        identifier=record['id'];data['id']=identifier
        assert client.get(ADMIN+kind+f'/query-by-id/{identifier}').json()['data']['id']==identifier
        assert client.post(ADMIN+kind+'/query-list',json={'pageNum':1,'pageSize':10}).json()['code']=='0000'
        assert client.put(ADMIN+kind+'/update-by-id',json=data).json()['data'] is True
        key={'admin-user':'user-id','ai-client-api':'api-id','ai-client-model':'model-id','ai-client-tool-mcp':'mcp-id'}[kind];field={'admin-user':'userId','ai-client-api':'apiId','ai-client-model':'modelId','ai-client-tool-mcp':'mcpId'}[kind]
        assert client.get(ADMIN+kind+'/query-by-'+key+'/'+data[field]).json()['data'][field]==data[field]
        assert client.put(ADMIN+kind+'/update-by-'+key,json=data).json()['code']=='0000'
        if kind!='admin-user':assert client.get(ADMIN+kind+'/query-enabled').json()['code']=='0000'
        if kind in ('admin-user','ai-client-tool-mcp'):assert client.get(ADMIN+kind+'/query-by-status/'+str(data['status'])).json()['code']=='0000'
    for kind,data in defs.items():
        key={'admin-user':'user-id','ai-client-api':'api-id','ai-client-model':'model-id','ai-client-tool-mcp':'mcp-id'}[kind]
        field={'admin-user':'userId','ai-client-api':'apiId','ai-client-model':'modelId','ai-client-tool-mcp':'mcpId'}[kind]
        for selector in ('id',key):
            copy={k:v for k,v in data.items() if k!='id'};copy[field]+='-delete-'+selector
            if kind=='admin-user':copy['username']+='-delete-'+selector
            assert client.post(ADMIN+kind+'/create',json=copy).json()['data'] is True
            record=client.get(ADMIN+kind+'/query-by-'+key+'/'+copy[field]).json()['data']
            assert client.delete(ADMIN+kind+'/delete-by-'+selector+'/'+str(record['id'] if selector=='id' else copy[field])).json()['data'] is True
    assert client.get(ADMIN+'admin-user/query-enabled').json()['code']=='0000'
    assert client.post(ADMIN+'ai-client-model/test/p5-model').json()['data']['ok']
    with httpx.Client(base_url=env[1],trust_env=False) as validation:
        assert validation.post(ADMIN+'admin-user/validate-login',json={'username':'p5-admin','password':'isolated-password'}).json()['data'] is True
        assert not validation.cookies.get('reactor_admin_token')
        assert validation.post(ADMIN+'admin-user/login',json={'username':'p5-admin','password':'wrong'}).json()['code']=='0002'
    assert client.get(ADMIN+'admin-user/query-by-username/secondary').json()['data']['userId']=='p5-secondary'
    assert client.get(ADMIN+'ai-client-model/query-by-api-id/p5-api').json()['data']
    assert client.get(ADMIN+'ai-client-model/query-by-model-type/chat').json()['data']
    assert client.get(ADMIN+'ai-client-tool-mcp/query-by-transport-type/stdio').json()['data']
    test=client.post(ADMIN+'ai-client-model/test-by-id/'+str(defs['ai-client-model']['id'])).json()['data'];assert test['ok'],test
    public=httpx.get(env[1]+ADMIN+'ai-client-model/query-enabled',trust_env=False).json()['data'];assert any(r['modelName']=='p5-actual-binding' for r in public)
    assert all('apiKey' not in r and 'password' not in r for r in public)
    frames=query(client,'p5-config','p5-config-run','配置即时生效',model='__db_binding__:'+str(defs['ai-client-model']['id']))
    assert frames[-1]['status']=='success'
    with env[3].connect() as c:
        assert c.execute(text("SELECT COUNT(*) FROM ai_agent_llm_invocation WHERE model_name='p1-fixed-upstream' AND run_id=(SELECT id FROM ai_agent_dialogue_run WHERE request_id='p5-config-run')")).scalar()>0
        assert any(r.get('model')=='p5-actual-binding' for r in httpx.get(os.environ['P1_TEST_MODEL_URL']+'/test/requests',trust_env=False).json())
        digest=c.execute(text('SELECT token_digest FROM ai_agent_admin_session ORDER BY id DESC LIMIT 1')).scalar();assert len(digest)==64 and digest!=client.cookies.get('reactor_admin_token')
    # Existing cookies lose permission immediately when the account/session changes.
    for state in ('disabled','expired','revoked'):
        login(env,client)
        with env[3].begin() as c:
            if state=='disabled':c.execute(text("UPDATE admin_user SET status=0 WHERE user_id='p5-admin'"))
            elif state=='expired':c.execute(text("UPDATE ai_agent_admin_session SET expires_at=DATE_SUB(NOW(),INTERVAL 1 SECOND) WHERE revoked_at IS NULL"))
            else:c.execute(text("UPDATE ai_agent_admin_session SET revoked_at=NOW() WHERE revoked_at IS NULL"))
        assert client.get(ADMIN+'admin-user/query-all').json()['code']=='0003',state
    login(env,client)
    bad=client.post(ADMIN+'ai-client-api/create',json=defs['ai-client-api'],headers={'Origin':'https://cross-site.invalid'});assert bad.json()['code']=='0003'
    sub=dict(agentKey='p5-reader',displayName='P5 阅读',whenToUse='Use P5 reader',systemPrompt='Read accurately',allowedTools=['workspace_read','workspace_read'],disallowedTools=[],maxSteps=3,status=1)
    assert client.post(ADMIN+'sub-agent-definitions/create',json=sub).json()['data'] is True
    assert client.get(ADMIN+'sub-agent-definitions/query-list').json()['data']
    assert client.get(ADMIN+'sub-agent-definitions/p5-reader').json()['data']['allowedTools']==['workspace_read']
    assert client.post(ADMIN+'sub-agent-definitions/reload').json()['data']==1
    assert 'emit_ui_tree' in client.get(ADMIN+'sub-agent-definitions/tool-catalog').json()['data']
    sub['status']=0;assert client.put(ADMIN+'sub-agent-definitions/update',json=sub).json()['data'] is True
    assert client.delete(ADMIN+'sub-agent-definitions/p5-reader').json()['data'] is True
    assert client.post(ADMIN+'sub-agent-definitions/create',json={**sub,'agentKey':'p5-delete-query'}).json()['data'] is True
    assert client.delete(ADMIN+'sub-agent-definitions/delete',params={'agentKey':'p5-delete-query'}).json()['data'] is True
    assert client.post(ADMIN+'sub-agent-definitions/create',json={**sub,'agentKey':'general-purpose'}).json()['code']=='0002'
    skill=dict(name='p5-legitimate',description='P5 real script',content='---\nname: p5-legitimate\ndescription: P5 real script\n---\nUse scripts deliberately.\n')
    assert client.post(ADMIN+'skills/create',json=skill).json()['code']=='0000'
    assert any(r['name']=='p5-legitimate' for r in client.get(ADMIN+'skills/list').json()['data'])
    assert client.post(ADMIN+'skills/reload').json()['data'] is True
    assert client.delete(ADMIN+'skills/p5-legitimate').json()['data'] is True
    packed=io.BytesIO()
    with zipfile.ZipFile(packed,'w') as z:z.writestr('p5-package/SKILL.md','---\nname: p5-package\ndescription: Real HTTP import\n---\nUse the package.\n')
    assert client.post(ADMIN+'skills/parse-package',files={'file':('p5.skill',packed.getvalue(),'application/zip')}).json()['code']=='0000'
    assert client.post(ADMIN+'skills/upload',files={'file':('p5.skill',packed.getvalue(),'application/zip')}).json()['code']=='0000'
    assert client.delete(ADMIN+'skills/p5-package').json()['data'] is True
    assert client.post(ADMIN+'skills/import-url',json={'url':os.environ['P5_TEST_IMAGE_URL']+'/p5/skill-package'}).json()['code']=='0000'
    assert client.delete(ADMIN+'skills/p5-url-package').json()['data'] is True
    visitor=httpx.Client(base_url=env[1],trust_env=False);visitor.get('/api/agent/visitor/bootstrap')
    for item in json.loads((Path(__file__).parents[1]/'reactor_backend/resources/admin_routes.json').read_text()):
        if item['path'].endswith(('/login','/validate-login')):continue
        path=__import__('re').sub(r'\{[^}]+\}','p5-denied',item['path'])
        denied=visitor.request(item['method'],path,json={});assert denied.json()['code']=='0003',item
    visitor.close()
    assert client.post(ADMIN+'admin-user/logout').json()['data'] is True
    assert client.get(ADMIN+'ai-client-api/query-all').json()['code']=='0003'
    login(env,client)


def test_p5_images_real_provider_wire_mask_batch_owner_history(env,client):
    from PIL import Image
    buf=io.BytesIO();Image.new('RGBA',(16,16),(255,0,0,0)).save(buf,format='PNG');ref='data:image/png;base64,'+base64.b64encode(buf.getvalue()).decode()
    batches=[dict(requestId='p5-images',prompt='生成测试图片',n=2,model='ignored-model'),dict(requestId='p5-edit',prompt='编辑原图',mode='edits',fileNames=[ref],maskFileNames=[ref],n=2),dict(requestId='p5-multi',prompt='融合参考图',fileNames=[ref,ref],n=5),dict(requestId='p5-fallback',prompt='fallback high resolution',size='2048x1152',n=4)]
    files=[]
    for body in batches:
        response=client.post('/api/agent/image-generation/generate',json=body,timeout=50);value=response.json();assert value['code']=='0000',value
        images=value['data'];files.extend(images['fileInfo']);assert images['requestId']==body['requestId'] and images['fileInfo']
        for file in images['fileInfo']:
            download=client.get(file['downloadUrl']);assert download.status_code==200 and Image.open(io.BytesIO(download.content)).size==(96,64)
        if 'fallback' in body['requestId']:assert images['usedFallback'] is True
    history=client.get('/api/agent/image-generation/history').json()['data'];assert history['total']==4 and all(b['images'] for b in history['list'])
    records=httpx.get(os.environ['P5_TEST_IMAGE_URL']+'/p5/requests',trust_env=False).json();assert all(r['body']['model']=='gpt-image-2.5-flare' for r in records)
    edit=next(r for r in records if r['path']=='edits' and r['body']['prompt']=='编辑原图');assert edit['files']['mask'] and edit['body']['n']=='2'
    multi=next(r for r in records if r['path']=='edits' and r['body']['prompt']=='融合参考图');assert len(multi['files']['image[]'])==2 and 'n' not in multi['body']
    assert any(r['path']=='chat' for r in records)
    with env[3].connect() as c:
        rows=c.execute(text("SELECT i.*,a.tool_call_id artifact_call FROM ai_agent_tool_output_image_generation i JOIN ai_agent_artifact a ON a.request_id=i.request_id AND a.tool_call_id=i.tool_call_id WHERE i.request_source='workspace' AND i.request_id LIKE 'p5-%' ")).mappings().all()
        assert rows and all(r['tool_call_id']==r['artifact_call'] and r['run_id'] is None for r in rows)
    with httpx.Client(base_url=env[1],trust_env=False) as other:
        other.get('/api/agent/visitor/bootstrap');assert other.get('/api/agent/image-generation/history').json()['data']['total']==0
        assert other.get(files[0]['downloadUrl']).status_code==403
        foreign=other.post('/api/agent/image-generation/generate',json={'requestId':'p5-foreign','prompt':'edit','fileNames':[files[0]['previewUrl']]});assert foreign.status_code==403
    assert client.post('/api/agent/image-generation/generate',json=batches[0]).json()['code']=='0002'
    assert client.post('/api/agent/image-generation/generate',json={'prompt':'invalid','size':'100x100'}).json()['code']=='0002'
    agent=query(client,'p5-image-agent','p5-image-agent-run','P5_IMAGE_AGENT',sessionFiles=[files[0]])
    assert agent[-1]['status']=='success'
    with env[3].connect() as c:
        stored=c.execute(text("SELECT i.*,a.file_name FROM ai_agent_tool_output_image_generation i JOIN ai_agent_artifact a ON a.tool_call_id=i.tool_call_id AND a.request_id=i.request_id WHERE i.request_source='agent' AND i.request_id='p5-image-agent-run'")).mappings().all()
        assert stored and all(r['source_image_count']==1 and r['run_id'] is not None and r['tool_invocation_id'] is not None for r in stored),stored
    records=httpx.get(os.environ['P5_TEST_IMAGE_URL']+'/p5/requests',trust_env=False).json()
    assert any(r['path']=='edits' and r['body']['prompt']=='P5 agent automatic attached reference' for r in records)
    (env[0]/'p5-image-evidence.json').write_text(json.dumps({'history':history,'providerRequests':records},ensure_ascii=False,indent=2))


def test_p5_genui_sse_export_history_featured_public_and_revocation(env,client):
    login(env,client);events=query(client,'p5-genui','p5-genui-run','P5_GENUI 生成可交互结果')
    assert events[-1]['status']=='success',events[-1]
    kinds=[e.get('resultMap',{}).get('eventData',{}).get('resultMap',{}).get('messageType') for e in events]
    assert {'ui_tree','ui_patch','html'}<=set(kinds),kinds
    history=client.get('/api/agent/conversation/sessions/p5-genui').json()['data'];assert 'ui_tree' in json.dumps(history)
    for format in ('pdf','docx'):
        response=client.post('/api/agent/genui/export/'+format,json={'tree':TREE});assert response.status_code==200,response.text[:100]
        (env[0]/('p5-genui.'+format)).write_bytes(response.content)
    # Publish an actual run's registered output. A sibling upload stays private.
    artifact=next(f for f in client.post('/v1/file_tool/get_file_list',json={'requestId':'p5-genui'}).json()['results'] if f['fileName'].endswith('.html'))
    private=client.post('/api/agent/file/upload',data={'sessionId':'p5-genui'},files={'file':('仅私有.txt',b'private-only','text/plain')});assert private.status_code==200
    body=dict(sessionId='p5-genui',title='P5 精选实际对话',summary='GenUI 与真实产物',coverUrl=artifact['domainUrl'],tags=['P5'],sortOrder=10,operator='forged-operator')
    assert client.post(ADMIN+'featured-conversations/create',json=body).json()['data'] is True
    fid='featured_p5-genui'
    assert httpx.get(env[1]+'/api/agent/featured-conversations/'+fid,trust_env=False).json()['data'] is None
    assert client.post(ADMIN+'featured-conversations/online/'+fid,params={'operator':'forged'}).json()['data'] is True
    public=httpx.get(env[1]+'/api/agent/featured-conversations/'+fid,trust_env=False).json()['data'];assert public['contentAvailable'] and 'ui_tree' in json.dumps(public['historyDetail'])
    assert httpx.get(artifact['domainUrl'],trust_env=False).status_code==200
    assert httpx.get(private.json()['previewUrl'],trust_env=False).status_code==403
    assert httpx.get(env[1]+'/api/agent/workspace/p5-genui/archive',trust_env=False).status_code==403
    with env[3].connect() as c:assert c.execute(text("SELECT updated_by FROM ai_agent_featured_conversation WHERE featured_id=:id"),{'id':fid}).scalar()=='p5-admin'
    assert client.post(ADMIN+'featured-conversations/query-list',json={'title':'P5','pageNo':1,'pageSize':20}).json()['data']['total']==1
    assert client.post(ADMIN+'featured-conversations/offline/'+fid,params={'operator':'x'}).json()['data'] is True
    assert httpx.get(artifact['domainUrl'],trust_env=False).status_code==403
    assert httpx.get(env[1]+'/api/agent/featured-conversations/'+fid,trust_env=False).json()['data'] is None
    archive=client.get('/api/agent/workspace/p5-genui/archive');assert archive.status_code==200
    with zipfile.ZipFile(io.BytesIO(archive.content)) as z:assert 'input/仅私有.txt' in z.namelist()
    mesh=next(f for f in client.post('/v1/file_tool/get_file_list',json={'requestId':'p5-genui'}).json()['results'] if f['fileName'].endswith('.gltf'))
    raw=client.get(mesh['domainUrl']);assert raw.headers['content-type'].startswith('model/gltf+json')
    ranged=client.get(mesh['domainUrl'],headers={'Range':'bytes=0-9'});assert ranged.status_code==206 and ranged.content==raw.content[:10]
    canvas=env[0]/'workspace/p5-genui/assets/canvas.html';original=canvas.read_bytes();canvas.unlink();canvas.symlink_to(env[0]/'connection.json')
    assert client.get(artifact['domainUrl']).status_code==403
    canvas.unlink();canvas.write_bytes(original)
    (env[0]/'p5-genui-events.json').write_text(json.dumps(events,ensure_ascii=False,indent=2))

async def test_p5_original_docgen_docread_dataprep_and_visitor_configuration(env,client):
    from reactor_backend.adapters.mysql import MySQLLedger
    from reactor_backend.adapters.tool_api import ToolAPI
    from reactor_backend.adapters.private_resources import PrivateResources
    from reactor_backend.bootstrap.config import Settings
    from reactor_backend.models.types import Query
    visitor=client.get('/api/agent/visitor/bootstrap').json()['data']['visitorId']
    sid='p5-documents';ledger=MySQLLedger(os.environ['P1_TEST_MYSQL_URL'])
    settings=Settings(os.environ['P1_TEST_MYSQL_URL'],env[2],env[0]/'workspace',tool_service_token=os.environ['P1_TEST_SERVICE_TOKEN'],private_tool_gateway=True)
    evidence=[]
    async with httpx.AsyncClient(trust_env=False,timeout=90) as http:
        api=ToolAPI(http,env[2],settings.workspace_root,90,settings=settings);resources=PrivateResources(ledger,api);api.resources=resources
        await resources.namespace(visitor,sid)
        q=Query(sid,sid,visitor,'P5 documents',sid,'user',0)
        async def tool(name,a):
            result=await api.execute(name,a,q);assert not result.failed,(name,result.error,result.structured_output)
            evidence.append({'tool':name,'arguments':a,'data':result.data,'files':result.files})
            return result
        generated={}
        try:
            for fmt in ('pdf','docx','html','markdown'):
                result=await tool('document_generate',dict(format=fmt,title='P5 原文档生成验收',content='# Reactor 迁移报告\n\n实际内容与导出测试\n\n| 项目 | 结果 |\n| --- | --- |\n| P5 | 完成 |',output_path='原报告.'+{'markdown':'md'}.get(fmt,fmt)))
                assert result.files;f=result.files[0];generated[fmt]=f
                r=client.get(f['domainUrl']);assert r.status_code==200 and len(r.content)>50
                (env[0]/('p5-docgen.'+{'markdown':'md'}.get(fmt,fmt))).write_bytes(r.content)
                if fmt=='pdf':
                    from pypdf import PdfReader
                    assert 'Reactor' in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(r.content)).pages)
                if fmt=='docx':
                    from docx import Document
                    assert '实际内容' in '\n'.join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
            result=await tool('slides_generate',dict(title='P5 Slide',slides=[{'title':'P5 real slide','body':'- Original PPTX renderer\n- 真实内容'}],output_path='slides.pptx'))
            assert result.files and zipfile.is_zipfile(io.BytesIO(client.get(result.files[0]['domainUrl']).content))
            result=await tool('excel_generator',dict(sheets=[{'name':'Sales','data':[['category','amount'],['A',10],['B',20]]}],output_path='sales.xlsx'));generated['excel']=result.files[0]
            from openpyxl import load_workbook
            assert load_workbook(io.BytesIO(client.get(result.files[0]['domainUrl']).content)).sheetnames==['Sales']
            await tool('checklist_generate',dict(items=[{'text':'验收文件'},{'text':'检查权限'}],title='P5 checklist',format='markdown',output_path='checklist.md'))
            await tool('template_filler',dict(template_string='Hello {{ name }}',variables={'name':'Reactor'},output_path='filled.txt'))
            await tool('chart_generator',dict(chart_type='bar',data={'categories':['A','B'],'series':[{'name':'Sales','values':[10,20]}]},title='P5 Chart',output_path='chart.png'))
            for action in ('create','get','list'):
                result=await tool('theme_designer',dict(action=action,name='p5-brand',primary='#336699'))
                assert any(key in result.data for key in ('saved','resolved','themes'))
            template=dict(action='save',name='p5-template',kind='document',content='# Hello {{ name }}',variables=[{'name':'name','required':True}])
            await tool('document_template',template)
            for action in ('get','list','preview','generate'):
                result=await tool('document_template',dict(action=action,name='p5-template',values={'name':'Original Template'},format='docx',output_path='templated.docx'))
                assert any(k in result.data for k in ('template','templates','rendered','produced_files'))
            # Read original generated files through positive per-session roots.
            for toolname,fmt in [('pdf_reader','pdf'),('pdf_structure','pdf'),('citation_extractor','pdf'),('word_reader','docx'),('html_processor','html'),('markdown_processor','markdown'),('excel_reader','excel')]:
                await tool(toolname,dict(file_path=generated[fmt]['fileName'],operation='read'))
            workspace=settings.workspace_root/sid;workspace.mkdir(exist_ok=True)
            (workspace/'source.csv').write_text('category,amount\nA,10\nA,20\nB,7.5\n')
            (workspace/'readme.txt').write_text('P5 Text processor actual content')
            await tool('text_processor',dict(file_path='readme.txt',operation='read'))
            await tool('csv_processor',dict(file_path='source.csv',operation='read'))
            records=[dict(category='A',amount=10),dict(category='A',amount=20),dict(category='B',amount=7.5)]
            await tool('data_clean',dict(data=records,operations=[{'type':'drop_duplicates'}]))
            await tool('data_aggregate',dict(data=records,operation='groupby',group_by=['category'],aggregations={'amount':'sum'}))
            await tool('data_transform',dict(data=records,transformations=[{'type':'rename','columns':{'amount':'value'}}]))
            await tool('data_merge',dict(left_data=records,right_data=[{'category':'A','label':'Alpha'},{'category':'B','label':'Beta'}],on=['category'],how='left'))
            await tool('data_validate',dict(data=records,schema={'required_columns':['category','amount']}))
            await tool('sql_query',dict(data=records,query='SELECT category,SUM(amount) total FROM data GROUP BY category'))
            for raw in (str(env[0]/'connection.json'),'../connection.json'):
                result=await api.execute('text_processor',dict(file_path=raw,operation='read'),q);assert result.failed
            other='p5-documents-other';await resources.namespace(visitor,other)
            foreign=settings.workspace_root/other;foreign.mkdir();(foreign/'private.txt').write_text('private')
            (workspace/'symlink.txt').symlink_to(foreign/'private.txt')
            (workspace/'includes').mkdir();(workspace/'includes/legitimate.txt').write_text('Hello {{ name }}')
            included=await tool('template_filler',dict(template_string='{% include "legitimate.txt" %}',include_paths=[str(workspace/'includes')],variables={'name':'Reactor'},output_path='include.txt'))
            assert included.data['rendered']=='Hello Reactor'
            (workspace/'includes/outside.txt').symlink_to(foreign/'private.txt')
            assert (await api.execute('template_filler',dict(template_string='{% include "outside.txt" %}',include_paths=[str(workspace/'includes')],output_path='deny.txt'),q)).failed
            assert (await api.execute('template_filler',dict(template_string="{{ ''.__class__.__mro__ }}",output_path='deny-attribute.txt'),q)).failed

            assert (await api.execute('text_processor',dict(file_path='symlink.txt',operation='read'),q)).failed
            with httpx.Client(base_url=env[1],trust_env=False) as another:
                another.get('/api/agent/visitor/bootstrap');visitor2=another.get('/api/agent/visitor/bootstrap').json()['data']['visitorId']
                q2=Query('p5-other-docgen','p5-other-docgen',visitor2,None,'trace','user',0)
                await resources.namespace(visitor2,q2.session_id)
                hidden=await api.execute('document_template',dict(action='list'),q2);assert not hidden.failed and all(t['name']!='p5-template' for t in hidden.data['templates'])
                themes=await api.execute('theme_designer',dict(action='list'),q2);assert all(t['name']!='p5-brand' for t in themes.data['themes'])
            await tool('document_template',dict(action='delete',name='p5-template'));await tool('theme_designer',dict(action='delete',name='p5-brand'))
            (env[0]/'p5-docgen-evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2))
        finally:await ledger.close()
